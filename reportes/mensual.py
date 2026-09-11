"""Excel mensual (base de la versión Plus): precios semanales, veredictos, aciertos y ahorro estimado
por galón si se hubieran seguido las alertas. Genera reportes/mensual-AAAA-MM.xlsx.

Ahorro por semana (regular, Q/galón):
  - Veredicto "Llena HOY" (alza) y el MEM subió:  ahorro = +cambio real (llenaste antes de la subida).
  - Veredicto "Espera" (baja) y el MEM bajó:       ahorro = |cambio real| (llenaste después de la bajada).
  - Veredicto contrario al movimiento real:        ahorro = -|cambio real| (costo de la señal fallida).
  - "Llena esta semana" o sin cambio:              0.
"""
from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from agente.comun import DATA, RAIZ, ahora_gt, leer_json, log  # noqa: E402

CARPETA = RAIZ / "reportes"


def ahorro_semana(predicho: str | None, cambio_real: float | None) -> float:
    if predicho is None or cambio_real is None:
        return 0.0
    if predicho == "alza":
        return round(cambio_real, 2) if cambio_real > 0 else round(-abs(cambio_real), 2) if cambio_real < 0 else 0.0
    if predicho == "baja":
        return round(abs(cambio_real), 2) if cambio_real < 0 else round(-abs(cambio_real), 2) if cambio_real > 0 else 0.0
    return 0.0


def _hoja(wb, titulo: str, encabezados: list[str], filas: list[list]):
    ws = wb.create_sheet(titulo)
    ws.append(encabezados)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F5FBF")
    for f in filas:
        ws.append(f)
    for i, h in enumerate(encabezados, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(14, min(60, len(h) + 4))
    ws.freeze_panes = "A2"
    return ws


def generar(mes: str | None = None) -> Path:
    hoy = ahora_gt().date()
    mes = mes or hoy.strftime("%Y-%m")
    precios = leer_json(DATA / "precios_historial.json", []) or []
    senales = leer_json(DATA / "senal_historial.json", []) or []
    aciertos = leer_json(DATA / "aciertos.json", {}) or {}

    wb = Workbook()
    wb.remove(wb.active)
    _hoja(wb, "Precios semanales", ["Fecha MEM", "Súper (Q/gal)", "Regular (Q/gal)", "Diésel (Q/gal)", "Fuente"],
          [[p["fecha"], p.get("superior"), p.get("regular"), p.get("diesel"), p.get("fuente")] for p in precios])
    _hoja(wb, "Veredictos diarios", ["Fecha", "Puntaje", "Tendencia", "Veredicto", "Cambio estimado (Q/gal)"],
          [[s["fecha"], s.get("puntaje"), s.get("tendencia"), s.get("veredicto"), s.get("cambio_estimado_q")] for s in senales])
    filas_ac, total_ahorro = [], 0.0
    for s in aciertos.get("semanas", []):
        a = ahorro_semana(s.get("predicho"), s.get("cambio_real_q"))
        total_ahorro += a
        filas_ac.append([s["fecha_martes"], s.get("predicho"), s.get("cambio_predicho_q"), s.get("real"), s.get("cambio_real_q"),
                         "Sí" if s.get("acierto") else "No" if s.get("acierto") is False else "Sin dato", a])
    ws = _hoja(wb, "Aciertos y ahorro", ["Martes", "Predicho", "Cambio predicho (Q/gal)", "Real", "Cambio real (Q/gal)", "Acierto", "Ahorro estimado (Q/gal)"], filas_ac)
    ws.append([])
    ws.append(["Total", "", "", "", "", f"{aciertos.get('porcentaje') or 0} % aciertos", round(total_ahorro, 2)])
    ws.append(["Ahorro por tanque de 12 galones", "", "", "", "", "", round(total_ahorro * 12, 2)])
    ws.append(["Ahorro para flotilla de 20 vehículos (12 gal c/u)", "", "", "", "", "", round(total_ahorro * 12 * 20, 2)])
    ws.append([])
    ws.append(["Generado", ahora_gt().isoformat(timespec="minutes")])

    CARPETA.mkdir(exist_ok=True)
    ruta = CARPETA / f"mensual-{mes}.xlsx"
    wb.save(ruta)
    log(f"Excel mensual generado: {ruta.name} ({len(precios)} semanas, {len(filas_ac)} martes, ahorro Q{total_ahorro:.2f}/gal)")
    return ruta


if __name__ == "__main__":  # pragma: no cover
    generar()
